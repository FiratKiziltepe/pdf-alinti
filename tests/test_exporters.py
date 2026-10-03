import base64
import csv
import io
import json
from dataclasses import replace

import pytest
import fitz
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook

from pdf_notes.exporters import (
    export_csv,
    export_docx,
    export_json,
    export_markdown,
    export_pdf,
    export_xlsx,
)
from pdf_notes.models import AnnotationRecord


@pytest.fixture
def annotation():
    return AnnotationRecord(
        id="örnek-1",
        filename="Çalışma.pdf",
        document_title="Araştırma ve öğrenme",
        page=3,
        kind="Highlight",
        color_name="Sarı",
        color_hex="#FFFF00",
        quote='İlk satır: “ölçüm”.\nİkinci satır; özgün alıntı.',
        context="Bu paragraf alıntıyı içerir ve daha geniş bağlam sağlar.",
        comment="Burayı tartışma bölümünde karşılaştır.",
        author="Çağrı",
        created="D:20260930123000Z",
        modified="D:20261001103000Z",
        rect=(10.0, 20.0, 300.0, 45.0),
        context_inferred=True,
    )


def test_json_losslessly_retains_text_metadata_and_geometry(annotation):
    data = json.loads(export_json([annotation], metadata={"açıklama": "Seçili notlar"}))
    assert data["version"] == 2
    row = data["records"][0]
    assert row["quote"] == annotation.quote
    assert row["context"] == annotation.context
    assert row["comment"] == annotation.comment
    assert "color_hex" not in row and "color_name" not in row
    assert row["author"] == "Çağrı"
    assert row["rect"] == [10.0, 20.0, 300.0, 45.0]
    assert row["context_inferred"] is True
    assert data["metadata"]["açıklama"] == "Seçili notlar"


def test_markdown_groups_documents_and_pages_with_separate_fields(annotation):
    first_page = replace(annotation, id="early", page=1)
    second_document = replace(annotation, id="other", filename="Başka.pdf", document_title="Diğer çalışma")
    content = export_markdown([annotation, second_document, first_page]).decode("utf-8")
    assert content.index("### Sayfa 1") < content.index("### Sayfa 3") < content.index("## Başka")
    # Markdown adds quote markers; removing them reconstructs the original text.
    assert "> " + annotation.quote.replace("\n", "\n> ") in content
    assert "> " + annotation.context in content
    assert annotation.comment in content
    assert "Sizin notunuz" in content and "background:#FFF1EB" in content
    assert "Not yazarı:** Çağrı" in content
    assert "FFFF00" not in content and "**Renk:" not in content
    assert "konumdan tahmin edildi" in content


def test_csv_bom_turkish_text_and_round_trip(annotation):
    content = export_csv([annotation])
    assert content.startswith(b"\xef\xbb\xbf")
    row = next(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    assert row["Alıntı"] == annotation.quote
    assert row["Bağlam"] == annotation.context
    assert row["Sizin notunuz"] == annotation.comment
    assert row["Not yazarı"] == "Çağrı"
    assert row["PDF dosyası"] == "Çalışma.pdf"
    assert "Renk kodu" not in row and "Renk" not in row
    assert row["Sayfa"] == "3"


@pytest.mark.parametrize("payload", ["=HYPERLINK(\"https://example.invalid\")", "+1", "-1", "@SUM(A1)", "   =1", "\t=1", "\r=1", "\x01text", "\x7ftext", "\u200b=1"])
def test_csv_protects_all_user_text_fields_from_formula_injection(annotation, payload):
    modified = replace(annotation, quote=payload, comment=payload, filename=payload, author=payload)
    row = next(csv.DictReader(io.StringIO(export_csv([modified]).decode("utf-8-sig"))))
    for key in ["Alıntı", "Sizin notunuz", "PDF dosyası", "Not yazarı"]:
        assert row[key] == "'" + payload


def test_word_is_valid_contains_grouped_annotations_and_color_shading(annotation):
    earlier = replace(annotation, page=1)
    document = Document(io.BytesIO(export_docx([annotation, earlier])))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert text.index("Sayfa 1") < text.index("Sayfa 3")
    assert annotation.quote in text
    assert annotation.context in text
    assert annotation.comment in text
    assert "Belge başlığı: Araştırma ve öğrenme" in text
    assert "Renk:" not in text and "#FFFF00" not in text
    assert "Not yazarı: Çağrı" in text
    quote_paragraph = next(p for p in document.paragraphs if p.text == annotation.quote)
    shading = quote_paragraph._p.find(".//" + qn("w:shd"))
    assert shading is not None
    assert shading.get(qn("w:fill")) == "FFFFBF"
    note = next(p for p in document.paragraphs if "Sizin notunuz" in p.text)
    assert annotation.comment in note.text
    assert note._p.find(".//" + qn("w:shd")).get(qn("w:fill")) == "FFF1EB"


def test_excel_is_valid_and_preserves_literal_formula_text(annotation):
    annotation = replace(annotation, quote="=SUM(A1:A2)", comment="   +1")
    workbook = load_workbook(io.BytesIO(export_xlsx([annotation])), data_only=False)
    worksheet = workbook.active
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:N2"
    assert worksheet["F2"].value == annotation.quote
    assert worksheet["F2"].data_type == "s"
    assert worksheet["H2"].value == annotation.comment
    assert worksheet["H2"].data_type == "s"
    assert worksheet["G2"].value == annotation.context
    assert worksheet["I2"].value == "Çağrı"
    assert worksheet["H2"].fill.fgColor.rgb.endswith("FFF1EB")
    assert worksheet["F2"].alignment.wrap_text is True


def test_excel_continues_long_text_without_silent_truncation(annotation):
    quote = "ö" * 32767 + "Son kısım"
    workbook = load_workbook(io.BytesIO(export_xlsx([replace(annotation, quote=quote)])))
    worksheet = workbook.active
    assert worksheet["F2"].value + worksheet["F3"].value == quote
    assert worksheet["N2"].value == "1/2"
    assert worksheet["N3"].value == "2/2"


def test_empty_exports_remain_valid_downloads():
    assert "not bulunamadı" in export_markdown([]).decode("utf-8")
    assert list(csv.DictReader(io.StringIO(export_csv([]).decode("utf-8-sig")))) == []
    assert json.loads(export_json([]))["records"] == []
    assert Document(io.BytesIO(export_docx([]))).paragraphs
    assert load_workbook(io.BytesIO(export_xlsx([]))).active.max_row == 1
    with fitz.open(stream=export_pdf([]), filetype="pdf") as pdf:
        assert len(pdf) == 1
        assert "not bulunamadı" in pdf[0].get_text()


def test_visual_note_image_survives_exports(annotation):
    with fitz.open() as doc:
        page = doc.new_page(width=100, height=200)
        page.draw_rect(fitz.Rect(10, 10, 90, 190), fill=(1, 0, 0))
        png = page.get_pixmap().tobytes("png")
    encoded = base64.b64encode(png).decode("ascii")
    record = replace(annotation, quote="", context="", image_base64=encoded)
    word = Document(io.BytesIO(export_docx([record])))
    assert len(word.inline_shapes) == 1
    shape = word.inline_shapes[0]
    image_part = word.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed]
    assert image_part.blob == png
    assert shape.height / shape.width == pytest.approx(2)
    assert all("seçili metin bulunamadı" not in p.text for p in word.paragraphs)
    workbook = load_workbook(io.BytesIO(export_xlsx([record])))
    visuals = workbook["Görsel notlar"]
    assert visuals["A2"].value == record.id
    assert visuals["B2"].value == record.filename
    assert visuals["C2"].value == record.page
    assert len(visuals._images) == 1
    assert visuals._images[0]._data() == png
    assert f"data:image/png;base64,{encoded}" in export_markdown([record]).decode("utf-8")
    assert json.loads(export_json([record]))["records"][0]["image_base64"] == encoded
    with fitz.open(stream=export_pdf([record]), filetype="pdf") as pdf:
        assert sum(len(page.get_images()) for page in pdf) >= 1


@pytest.mark.parametrize("include_context", [True, False])
def test_context_option_in_every_export(annotation, include_context):
    record = replace(annotation, context="UNIQUE_CONTEXT_ONLY", comment="UNIQUE_USER_NOTE")
    md = export_markdown([record], include_context=include_context).decode("utf-8")
    docx = Document(io.BytesIO(export_docx([record], include_context=include_context)))
    word_text = "\n".join(p.text for p in docx.paragraphs)
    row = next(csv.DictReader(io.StringIO(export_csv([record], include_context=include_context).decode("utf-8-sig"))))
    sheet = load_workbook(io.BytesIO(export_xlsx([record], include_context=include_context))).active
    excel_text = str(list(sheet.values))
    json_data = json.loads(export_json([record], include_context=include_context))["records"][0]
    with fitz.open(stream=export_pdf([record], include_context=include_context), filetype="pdf") as pdf:
        pdf_text = "\n".join(p.get_text() for p in pdf)
    for text in (md, word_text, str(row), excel_text, str(json_data), pdf_text):
        assert (record.context in text) is include_context
        assert record.comment in text
        assert "#FFFF00" not in text
    assert ("Bağlam" in row) is include_context
    assert ("context" in json_data) is include_context
    assert ("context_inferred" in json_data) is include_context
    assert record.context == "UNIQUE_CONTEXT_ONLY"  # exports never mutate the workspace


def test_pdf_unicode_colored_notes_and_long_content(annotation):
    record = replace(annotation, comment=("Türkçe: İı Şş Ğğ Çç Öö Üü <etiket> & not.\n" * 110) + "SON_NOT", quote="ALINTI_BASLANGIC γ α β × ﬀ ﬁ")
    with fitz.open(stream=export_pdf([record]), filetype="pdf") as pdf:
        assert len(pdf) > 1
        text = "\n".join(page.get_text() for page in pdf)
        assert "ALINTI_BASLANGIC" in text and "SON_NOT" in text
        assert "γ α β × ﬀ ﬁ" in text
        assert "Türkçe: İı Şş Ğğ Çç Öö Üü <etiket> & not." in text
        assert "Sizin notunuz" in text
        assert "Renk:" not in text and "#FFFF00" not in text
        fills = [drawing["fill"] for page in pdf for drawing in page.get_drawings() if drawing["fill"]]
        assert any(all(abs(actual - expected) < 0.01 for actual, expected in zip(fill, (1, 241 / 255, 235 / 255))) for fill in fills)
