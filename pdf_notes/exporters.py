"""Downloadable, Unicode-safe exports of the selected PDF annotations.

PDF text is exported as supplied by the extractor: no citation metadata is
guessed. CSV protects spreadsheet readers from interpreting annotation text
as formulas; XLSX writes every text value as a literal string.
"""

from __future__ import annotations

import base64
import csv
import io
import html
import json
import re
import unicodedata
from collections import OrderedDict
from dataclasses import asdict
from typing import Any, Mapping

from .models import AnnotationRecord
from .pdf_export import export_pdf


_COLUMNS = (
    ("id", "Kimlik"),
    ("filename", "PDF dosyası"),
    ("document_title", "Belge başlığı"),
    ("page", "Sayfa"),
    ("kind", "İşaret türü"),
    ("quote", "Alıntı"),
    ("context", "Bağlam"),
    ("comment", "Sizin notunuz"),
    ("author", "Not yazarı"),
    ("created", "Oluşturulma tarihi"),
    ("modified", "Değiştirilme tarihi"),
    ("context_inferred", "Bağlam konumdan tahmin edildi"),
    ("rect", "PDF koordinatları"),
)
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_HEX_COLOR = re.compile(r"^#?([0-9a-fA-F]{6})$")


def _group_records(records: list[AnnotationRecord]):
    """Keep document order and put each document's pages in reading order."""
    documents: OrderedDict[tuple[str, str], dict[int, list[AnnotationRecord]]] = OrderedDict()
    for record in records:
        pages = documents.setdefault((record.filename, record.document_title), {})
        pages.setdefault(record.page, []).append(record)
    for (filename, title), pages in documents.items():
        yield filename, title, [(page, pages[page]) for page in sorted(pages)]


def _markdown_label(value: str) -> str:
    """Prevent document metadata from turning into Markdown structure."""
    value = value.replace("\r", " ").replace("\n", " ")
    return re.sub(r"([\\`*_{}\[\]()#+.!|<>])", r"\\\1", value)


def _blockquote(value: str) -> str:
    # Splitting on literal newlines keeps even trailing and empty quoted lines.
    return "\n".join("> " + line for line in value.split("\n"))


def _rect_text(rect: tuple[float, float, float, float] | None) -> str:
    return json.dumps(rect, ensure_ascii=False) if rect is not None else ""


def _csv_literal(value: Any) -> Any:
    """Neutralize formulas, including formulas hidden behind leading spaces."""
    if not isinstance(value, str) or not value:
        return value
    offset = 0
    has_control = False
    while offset < len(value):
        character = value[offset]
        control = unicodedata.category(character) in ("Cc", "Cf")
        if not character.isspace() and not control:
            break
        has_control = has_control or control
        offset += 1
    if value[offset:].startswith(("=", "+", "-", "@")) or has_control:
        return "'" + value
    return value


def _xml_text(value: str) -> str:
    """Represent XML-forbidden control characters instead of failing a download."""
    return _INVALID_XML.sub(lambda match: f"\\u{ord(match[0]):04x}", value)


def _columns(include_context: bool):
    return tuple(column for column in _COLUMNS if include_context or column[0] not in {"context", "context_inferred"})


def _row(record: AnnotationRecord, include_context: bool) -> list[Any]:
    values = asdict(record)
    values["rect"] = _rect_text(record.rect)
    values["context_inferred"] = "Evet" if record.context_inferred else "Hayır"
    return [values[key] for key, _label in _columns(include_context)]


def export_markdown(records: list[AnnotationRecord], *, include_context: bool = True) -> bytes:
    """Return a Turkish Markdown report grouped by document and PDF page."""
    lines = ["# PDF alıntıları ve notları", ""]
    if not records:
        lines.extend(["Dışa aktarılacak not bulunamadı.", ""])
    for filename, title, pages in _group_records(records):
        lines.extend([f"## {_markdown_label(filename)}", ""])
        if title:
            lines.extend([f"**Belge başlığı:** {_markdown_label(title)}", ""])
        for page, page_records in pages:
            lines.extend([f"### Sayfa {page}", ""])
            for index, record in enumerate(page_records, start=1):
                lines.extend([
                    f"#### {index}. {_markdown_label(record.kind)}", "",
                ])
                if record.author:
                    lines.extend([f"**Not yazarı:** {_markdown_label(record.author)}", ""])
                if record.quote:
                    lines.extend(["**Alıntı:**", "", _blockquote(record.quote), ""])
                elif not record.image_base64:
                    lines.extend(["*Bu işaret için seçili metin bulunamadı.*", ""])
                if record.image_base64:
                    lines.extend([f"![Çerçeve içindeki alan · Sayfa {record.page}](data:image/png;base64,{record.image_base64})", ""])
                if include_context and record.context:
                    suffix = " (konumdan tahmin edildi)" if record.context_inferred else ""
                    lines.extend([f"**Bağlam{suffix}:**", "", _blockquote(record.context), ""])
                if record.comment:
                    note = html.escape(record.comment).replace("\n", "<br>")
                    lines.extend([f'<div style="background:#FFF1EB;color:#805040;padding:14px;border-left:3px solid #F06457;border-radius:6px"><strong>Sizin notunuz</strong><br>{note}</div>', ""])
                if record.created:
                    lines.extend([f"**Oluşturulma tarihi:** {_markdown_label(record.created)}", ""])
                if record.modified:
                    lines.extend([f"**Değiştirilme tarihi:** {_markdown_label(record.modified)}", ""])
                lines.extend(["---", ""])
    return "\n".join(lines).encode("utf-8")


def export_csv(records: list[AnnotationRecord], *, include_context: bool = True) -> bytes:
    """Return CSV with a BOM so spreadsheet applications detect Turkish UTF-8."""
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow([label for _key, label in _columns(include_context)])
    for record in records:
        writer.writerow([_csv_literal(value) for value in _row(record, include_context)])
    return stream.getvalue().encode("utf-8-sig")


def export_json(
    records: list[AnnotationRecord], metadata: Mapping[str, Any] | None = None,
    *, include_context: bool = True,
) -> bytes:
    """Return selected fields and images, without color metadata or hidden context."""
    excluded = {"color_name", "color_hex"}
    if not include_context:
        excluded.update({"context", "context_inferred"})
    rows = [{key: value for key, value in asdict(record).items() if key not in excluded} for record in records]
    payload: dict[str, Any] = {"version": 2, "records": rows}
    if metadata is not None:
        payload["metadata"] = dict(metadata)
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def export_docx(records: list[AnnotationRecord], *, include_context: bool = True) -> bytes:
    """Return a Word report with quotation shading matching annotation colors."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    document.core_properties.title = "PDF alıntıları ve notları"
    document.core_properties.language = "tr-TR"
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("PDF alıntıları ve notları", 0)
    if not records:
        document.add_paragraph("Dışa aktarılacak not bulunamadı.")

    def label_paragraph(label: str, value: str):
        paragraph = document.add_paragraph()
        paragraph.add_run(label + ": ").bold = True
        paragraph.add_run(_xml_text(value))
        return paragraph

    for filename, title, pages in _group_records(records):
        document.add_heading(_xml_text(filename), level=1)
        if title:
            label_paragraph("Belge başlığı", title)
        for page, page_records in pages:
            document.add_heading(f"Sayfa {page}", level=2)
            for index, record in enumerate(page_records, start=1):
                document.add_heading(_xml_text(f"{index}. {record.kind}"), level=3)
                if record.author:
                    label_paragraph("Not yazarı", record.author)
                if record.quote:
                    document.add_paragraph("Alıntı:").runs[0].bold = True
                    paragraph = document.add_paragraph(_xml_text(record.quote))
                    paragraph.paragraph_format.left_indent = Inches(0.2)
                    match = _HEX_COLOR.fullmatch(record.color_hex)
                    hex_color = match[1] if match else "EDEDED"
                    # Pale shading preserves legibility even for red/black marks.
                    pastel = "".join(f"{round(int(hex_color[i:i + 2], 16) * 0.25 + 255 * 0.75):02X}" for i in (0, 2, 4))
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), pastel)
                    paragraph._p.get_or_add_pPr().append(shading)
                elif not record.image_base64:
                    document.add_paragraph("Bu işaret için seçili metin bulunamadı.")
                if record.image_base64:
                    document.add_paragraph("Alıntı · Çerçeve içindeki alan:").runs[0].bold = True
                    picture = document.add_picture(io.BytesIO(base64.b64decode(record.image_base64)))
                    # Fit both wide tables and tall page selections on the page.
                    scale = min(Inches(6) / picture.width, Inches(7.5) / picture.height)
                    picture.width = round(picture.width * scale)
                    picture.height = round(picture.height * scale)
                if include_context and record.context:
                    label = "Bağlam (konumdan tahmin edildi)" if record.context_inferred else "Bağlam"
                    label_paragraph(label, record.context)
                if record.comment:
                    paragraph = document.add_paragraph()
                    paragraph.add_run("Sizin notunuz\n").bold = True
                    paragraph.add_run(_xml_text(record.comment))
                    paragraph.paragraph_format.space_before = Pt(10)
                    paragraph.paragraph_format.space_after = Pt(12)
                    paragraph.paragraph_format.left_indent = Inches(0.15)
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "FFF1EB")
                    paragraph._p.get_or_add_pPr().append(shading)
                    borders = OxmlElement("w:pBdr")
                    for side in ("top", "left", "bottom", "right"):
                        border = OxmlElement(f"w:{side}")
                        for key, value in {"val": "single", "sz": "6", "space": "8", "color": "FFE1D4"}.items():
                            border.set(qn(f"w:{key}"), value)
                        borders.append(border)
                    paragraph._p.get_or_add_pPr().append(borders)
                if record.created:
                    label_paragraph("Oluşturulma tarihi", record.created)
                if record.modified:
                    label_paragraph("Değiştirilme tarihi", record.modified)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def export_xlsx(records: list[AnnotationRecord], *, include_context: bool = True) -> bytes:
    """Return a formatted spreadsheet containing literal strings only.

    Excel limits individual cells to 32,767 characters. Longer values are
    continued onto adjacent rows with a "Metin parçası" sequence number so
    text is not silently truncated. JSON retains the original unsplit fields.
    """
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "PDF notları"
    columns = _columns(include_context)
    headers = [label for _key, label in columns] + ["Metin parçası"]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24495C")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[1].height = 32
    row_number = 2
    for record in records:
        values = [_xml_text(value) if isinstance(value, str) else value for value in _row(record, include_context)]
        parts = max([1] + [(len(value) + 32766) // 32767 for value in values if isinstance(value, str)])
        for part in range(parts):
            for column, value in enumerate(values, start=1):
                # Repeat ordinary metadata but keep long values as numbered chunks.
                if isinstance(value, str) and len(value) > 32767:
                    value = value[part * 32767:(part + 1) * 32767]
                cell = worksheet.cell(row_number, column, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            worksheet.cell(row_number, len(headers), f"{part + 1}/{parts}").data_type = "s"
            comment_column = next(index for index, (key, _) in enumerate(columns, start=1) if key == "comment")
            cell = worksheet.cell(row_number, comment_column)
            cell.fill = PatternFill("solid", fgColor="FFF1EB")
            cell.font = Font(color="805040")
            row_number += 1
    widths = [{"quote": 65, "context": 65, "comment": 55, "page": 8}.get(key, 26) for key, _ in columns] + [16]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    image_records = [record for record in records if record.image_base64]
    if image_records:
        images = workbook.create_sheet("Görsel notlar")
        images.append(["Kimlik", "PDF dosyası", "Sayfa", "Sizin notunuz", "Çerçeve içindeki alan"])
        for cell in images[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="6D28D9")
        for column, width in zip("ABCDE", (34, 30, 10, 45, 85)):
            images.column_dimensions[column].width = width
        for row, record in enumerate(image_records, start=2):
            # The full comment remains in the main sheet, including long-text continuations.
            for column, value in enumerate((record.id, record.filename, record.page, record.comment[:32767]), start=1):
                cell = images.cell(row, column, _xml_text(value) if isinstance(value, str) else value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            images.cell(row, 4).fill = PatternFill("solid", fgColor="FFF1EB")
            images.cell(row, 4).font = Font(color="805040")
            picture = Image(io.BytesIO(base64.b64decode(record.image_base64)))
            scale = min(1, 580 / picture.width, 520 / picture.height)
            picture.width = round(picture.width * scale)
            picture.height = round(picture.height * scale)
            images.add_image(picture, f"E{row}")
            images.row_dimensions[row].height = max(55, (picture.height + 16) * 0.75)
        images.freeze_panes = "E2"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
